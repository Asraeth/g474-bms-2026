"""
Checks the firmware's voltage -> SoC model against "Battery Tracker.xlsx" itself: reads the Chebyshev
coefficients and the voltage range from the sheet "Model" and evaluates the VToSoc formula the way
Excel does, COS(k * ACOS(t)), over the curve the soc_model test wrote.

    python test/check_soc.py [test/build/soc_model.csv] ["Supporting Docs/Battery Tracker.xlsx"]
"""

import csv
import math
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def main():
    curve = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "test", "build", "soc_model.csv")
    book = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, "Supporting Docs", "Battery Tracker.xlsx")

    sheet = openpyxl.load_workbook(book, data_only=True)["Model"]
    vmin, vmax = sheet["B1"].value, sheet["B2"].value
    coeffs = []
    for row in range(6, 40):
        k, d = sheet.cell(row, 1).value, sheet.cell(row, 2).value
        if k is None or d is None:
            break
        assert k == len(coeffs), f"coefficient order in Model!A{row}"
        coeffs.append(d)

    def vtosoc(v):
        t = max(-1.0, min(1.0, 2 * (v - vmin) / (vmax - vmin) - 1))
        return 100 * sum(d * math.cos(k * math.acos(t)) for k, d in enumerate(coeffs))

    worst = 0.0
    points = 0
    with open(curve, newline="") as f:
        for v, fw in csv.reader(f):
            v, fw = float(v), float(fw)
            ref = min(100.0, max(0.0, vtosoc(v)))      # The firmware clamps to 0..100
            worst = max(worst, abs(fw - ref))
            points += 1

    ok = worst < 0.01
    print(f"SoC model: {len(coeffs)} coefficients from {os.path.basename(book)}, {points} points, "
          f"largest difference {worst:.5f} % -> {'PASS' if ok else 'FAIL'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
